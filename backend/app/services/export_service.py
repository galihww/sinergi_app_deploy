from __future__ import annotations

import re
import unicodedata
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

import mistune
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


LEGAL_HEADINGS = {
    "IDENTITAS TERDAKWA",
    "DUDUK PERKARA",
    "TENTANG DUDUK PERKARA",
    "DAKWAAN",
    "TUNTUTAN PIDANA",
    "PERTIMBANGAN HUKUM",
    "MENIMBANG",
    "MENGINGAT",
    "MENGADILI",
    "MEMUTUSKAN",
    "AMAR PUTUSAN",
}
PAGE_MARKER = re.compile(r"^\s*\[HALAMAN\s+(\d+)\]\s*$", re.IGNORECASE)


def _exported_at() -> str:
    return datetime.now(timezone.utc).astimezone().strftime("%d %B %Y, %H:%M %Z")


def safe_filename(value: str, fallback: str = "legal-verse-export") -> str:
    normalized = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    normalized = re.sub(r"[^A-Za-z0-9._-]+", "-", normalized).strip(".-_")
    normalized = re.sub(r"-{2,}", "-", normalized)
    return normalized[:120] or fallback


def _escape_markdown(value: Any) -> str:
    return str(value or "").replace("\\", "\\\\").replace("[", "\\[").replace("]", "\\]")


def _source_markdown(sources: Iterable[dict[str, Any]]) -> str:
    rows: list[str] = []
    for index, source in enumerate(sources, start=1):
        title = source.get("title") or source.get("label") or source.get("id") or f"Sumber {index}"
        rows.append(f"### {index}. {title}")
        if source.get("score") is not None:
            try:
                rows.append(f"Skor relevansi: {float(source['score']):.3f}")
            except (TypeError, ValueError):
                rows.append(f"Skor relevansi: {source['score']}")
        if source.get("reason"):
            rows.append(f"Alasan: {_escape_markdown(source['reason'])}")
        if source.get("excerpt") or source.get("text"):
            rows.extend(["", str(source.get("excerpt") or source.get("text")).strip()])
        rows.append("")
    return "\n".join(rows).rstrip()


def build_pdf_markdown(
    name: str,
    text: str,
    *,
    page_count: int = 0,
    ocr_used: bool = False,
    warnings: Iterable[str] = (),
) -> tuple[str, str]:
    title = (Path(name).stem or "Putusan").replace("\n", " ").replace("\r", " ").strip()
    metadata = [
        f"**Dokumen sumber:** {_escape_markdown(name)}",
        f"**Diekspor:** {_exported_at()}",
    ]
    if page_count:
        metadata.append(f"**Jumlah halaman:** {page_count}")
    metadata.append(f"**OCR digunakan:** {'Ya' if ocr_used else 'Tidak'}")

    body: list[str] = []
    for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        marker = PAGE_MARKER.match(line)
        stripped = line.strip()
        if marker:
            body.extend(["", f"## Halaman {marker.group(1)}", ""])
        elif stripped.upper().rstrip(":") in LEGAL_HEADINGS and len(stripped) <= 80:
            body.extend(["", f"### {stripped.rstrip(':').title()}", ""])
        else:
            body.append(line.rstrip())

    warning_lines = [f"- {_escape_markdown(item)}" for item in warnings if str(item).strip()]
    warning_section = f"\n\n## Catatan Ekstraksi\n\n{chr(10).join(warning_lines)}" if warning_lines else ""
    markdown = (
        f"# {title}\n\n"
        + "  \n".join(metadata)
        + "\n\n> Konversi ini bukan salinan resmi. Dokumen PDF sumber tetap menjadi dokumen acuan.\n\n"
        + "\n".join(body).strip()
        + warning_section
        + "\n"
    )
    return markdown, safe_filename(title)


def _message_content(message: dict[str, Any]) -> str:
    return str(message.get("content") or "").strip()


def build_chat_message_markdown(
    title: str,
    messages: list[dict[str, Any]],
    message_id: str,
) -> tuple[str, str]:
    title = (title or "Jawaban Legal-Verse").replace("\n", " ").replace("\r", " ").strip()
    selected_index = next((i for i, item in enumerate(messages) if item.get("id") == message_id), -1)
    if selected_index < 0:
        raise ValueError("Pesan tidak ditemukan.")
    selected = messages[selected_index]
    if selected.get("role") != "assistant" or not _message_content(selected):
        raise ValueError("Hanya jawaban asisten yang dapat diekspor.")
    question = next(
        (_message_content(messages[i]) for i in range(selected_index - 1, -1, -1) if messages[i].get("role") == "user"),
        "",
    )
    sections = [
        f"# {title}",
        "",
        f"**Diekspor:** {_exported_at()}",
        "",
    ]
    if question:
        sections.extend(["## Pertanyaan", "", question, ""])
    sections.extend(["## Jawaban", "", _message_content(selected)])
    return "\n".join(sections).rstrip() + "\n", safe_filename(f"{title}-jawaban")


def build_rag_sources_markdown(
    title: str,
    messages: list[dict[str, Any]],
    message_id: str,
) -> tuple[str, str]:
    title = (title or "Hasil Retrieval RAG").replace("\n", " ").replace("\r", " ").strip()
    selected_index = next((i for i, item in enumerate(messages) if item.get("id") == message_id), -1)
    if selected_index < 0:
        raise ValueError("Pesan tidak ditemukan.")
    selected = messages[selected_index]
    if selected.get("role") != "assistant":
        raise ValueError("Hasil RAG hanya tersedia untuk jawaban asisten.")
    sources = selected.get("sources") or []
    if not sources:
        raise ValueError("Pesan ini tidak memiliki hasil retrieval RAG.")
    question = next(
        (_message_content(messages[i]) for i in range(selected_index - 1, -1, -1) if messages[i].get("role") == "user"),
        "",
    )
    sections = [
        f"# Hasil Retrieval RAG — {title}",
        "",
        f"**Diekspor:** {_exported_at()}",
        f"**Jumlah potongan:** {len(sources)}",
        "",
    ]
    if question:
        sections.extend(["## Pertanyaan Retrieval", "", question, ""])
    sections.extend(["## Potongan Sumber RAG", "", _source_markdown(sources)])
    return "\n".join(sections).rstrip() + "\n", safe_filename(f"{title}-hasil-rag")


def build_chat_session_markdown(
    title: str,
    messages: list[dict[str, Any]],
) -> tuple[str, str]:
    title = (title or "Percakapan Legal-Verse").replace("\n", " ").replace("\r", " ").strip()
    sections = [
        f"# {title}",
        "",
        f"**Diekspor:** {_exported_at()}",
        "",
        "## Percakapan",
    ]
    exported = 0
    for message in messages:
        content = _message_content(message)
        if not content or message.get("role") not in {"user", "assistant"}:
            continue
        exported += 1
        label = "Pengguna" if message["role"] == "user" else "Legal-Verse AI"
        sections.extend(["", f"### {label}", "", content])
    if not exported:
        raise ValueError("Percakapan belum memiliki pesan yang dapat diekspor.")
    return "\n".join(sections).rstrip() + "\n", safe_filename(title)


def markdown_bytes(markdown: str) -> bytes:
    return markdown.rstrip().encode("utf-8") + b"\n"


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margin(cell, top: int = 80, start: int = 100, bottom: int = 80, end: int = 100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _add_inline(paragraph, tokens: list[dict[str, Any]], *, bold: bool = False, italic: bool = False) -> None:
    for token in tokens:
        kind = token.get("type")
        if kind in {"text", "inline_html"}:
            run = paragraph.add_run(str(token.get("raw") or ""))
            run.bold, run.italic = bold, italic
        elif kind in {"strong", "emphasis", "strikethrough", "link"}:
            _add_inline(
                paragraph,
                token.get("children") or [],
                bold=bold or kind == "strong",
                italic=italic or kind == "emphasis",
            )
            if kind == "link" and token.get("attrs", {}).get("url"):
                visible = "".join(str(child.get("raw") or "") for child in token.get("children") or [])
                url = str(token["attrs"]["url"])
                if url and url != visible:
                    paragraph.add_run(f" ({url})")
        elif kind == "codespan":
            run = paragraph.add_run(str(token.get("raw") or ""))
            run.font.name = "Courier New"
            run.font.size = Pt(9.5)
        elif kind in {"softbreak", "linebreak"}:
            paragraph.add_run().add_break(WD_BREAK.LINE)
        elif token.get("children"):
            _add_inline(paragraph, token["children"], bold=bold, italic=italic)


def _inline_tokens(token: dict[str, Any]) -> list[dict[str, Any]]:
    return token.get("children") or [{"type": "text", "raw": token.get("raw", "")}]


def _render_blocks(document: Document, tokens: list[dict[str, Any]], *, list_level: int = 0) -> None:
    for token in tokens:
        kind = token.get("type")
        if kind == "blank_line":
            continue
        if kind == "heading":
            level = max(1, min(4, int(token.get("attrs", {}).get("level", 1))))
            paragraph = document.add_paragraph(style=f"Heading {level}")
            _add_inline(paragraph, _inline_tokens(token))
        elif kind == "paragraph":
            paragraph = document.add_paragraph()
            _add_inline(paragraph, _inline_tokens(token))
        elif kind == "block_quote":
            for child in token.get("children") or []:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.35)
                paragraph.paragraph_format.right_indent = Inches(0.2)
                paragraph.paragraph_format.space_after = Pt(6)
                _add_inline(paragraph, _inline_tokens(child), italic=True)
        elif kind == "block_code":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            run = paragraph.add_run(str(token.get("raw") or "").rstrip())
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        elif kind == "list":
            ordered = bool(token.get("attrs", {}).get("ordered"))
            style = "List Number" if ordered else "List Bullet"
            for item in token.get("children") or []:
                children = item.get("children") or []
                first = True
                for child in children:
                    if child.get("type") in {"block_text", "paragraph"}:
                        paragraph = document.add_paragraph(style=style if first else None)
                        paragraph.paragraph_format.left_indent = Inches(0.25 + list_level * 0.2)
                        _add_inline(paragraph, _inline_tokens(child))
                        first = False
                    elif child.get("type") == "list":
                        _render_blocks(document, [child], list_level=list_level + 1)
        elif kind == "table":
            head = next((x for x in token.get("children") or [] if x.get("type") == "table_head"), None)
            body = next((x for x in token.get("children") or [] if x.get("type") == "table_body"), None)
            header_cells = head.get("children") if head else []
            rows = body.get("children") if body else []
            column_count = len(header_cells) or max((len(row.get("children") or []) for row in rows), default=1)
            table = document.add_table(rows=1 if header_cells else 0, cols=column_count)
            table.style = "Table Grid"
            table.autofit = True
            if header_cells:
                for index, cell_token in enumerate(header_cells):
                    cell = table.rows[0].cells[index]
                    _shade_cell(cell, "3F3F46")
                    _set_cell_margin(cell)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    _add_inline(cell.paragraphs[0], _inline_tokens(cell_token), bold=True)
                    for run in cell.paragraphs[0].runs:
                        run.font.color.rgb = RGBColor(255, 255, 255)
            for row_token in rows:
                cells = table.add_row().cells
                for index, cell_token in enumerate(row_token.get("children") or []):
                    _set_cell_margin(cells[index])
                    _add_inline(cells[index].paragraphs[0], _inline_tokens(cell_token))
        elif kind == "thematic_break":
            document.add_paragraph()
        elif token.get("children"):
            _render_blocks(document, token["children"], list_level=list_level)


def docx_bytes(title: str, markdown: str) -> bytes:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(24, 24, 27)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.line_spacing = 1.08
    for level, size in ((1, 20), (2, 15), (3, 12.5), (4, 11)):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(24, 24, 27)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if level > 1 else 0)
        style.paragraph_format.space_after = Pt(6)

    document.core_properties.title = title
    document.core_properties.author = "Legal-Verse"
    parser = mistune.create_markdown(renderer="ast", plugins=["table", "strikethrough", "task_lists", "url"])
    tokens = parser(markdown)
    _render_blocks(document, tokens)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
