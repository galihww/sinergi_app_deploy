from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
from unittest import mock

import pytest
from fastapi import HTTPException
from docx import Document

from app.services.export_service import (
    build_chat_message_markdown,
    build_chat_session_markdown,
    build_pdf_markdown,
    docx_bytes,
    markdown_bytes,
    safe_filename,
)
from app.controllers.export_controller import export_document
from app.schemas import ExportRequest


MESSAGES = [
    {"id": "u1", "role": "user", "content": "Apa amar putusan perkara ini?"},
    {
        "id": "a1",
        "role": "assistant",
        "content": "## Ringkasan\n\nTerdakwa dinyatakan bersalah.\n\n| Unsur | Temuan |\n|---|---|\n| Niat | Terbukti |",
        "thinking": "Rahasia penalaran yang tidak boleh diekspor",
        "sources": [
            {
                "title": "Putusan 123/Pid.B/2025/PN Jkt",
                "score": 0.91234,
                "excerpt": "Menyatakan terdakwa terbukti secara sah dan meyakinkan.",
            }
        ],
    },
]


def test_pdf_markdown_preserves_page_boundaries_and_warnings() -> None:
    markdown, stem = build_pdf_markdown(
        "Putusan Nomor 7.pdf",
        "[HALAMAN 1]\nMENGADILI\nPidana dua tahun.\n[HALAMAN 2]\nSelesai.",
        page_count=2,
        ocr_used=True,
        warnings=["Halaman 2 menggunakan OCR."],
    )

    assert stem == "Putusan-Nomor-7"
    assert "## Halaman 1" in markdown
    assert "### Mengadili" in markdown
    assert "**OCR digunakan:** Ya" in markdown
    assert "Halaman 2 menggunakan OCR." in markdown
    assert "bukan salinan resmi" in markdown


def test_single_answer_export_includes_question_and_sources_but_not_thinking() -> None:
    markdown, stem = build_chat_message_markdown("Analisis Putusan", MESSAGES, "a1")

    assert stem == "Analisis-Putusan-jawaban"
    assert "## Pertanyaan" in markdown
    assert "Apa amar putusan" in markdown
    assert "## Jawaban" in markdown
    assert "## Sumber Jawaban" in markdown
    assert "Skor relevansi: 0.912" in markdown
    assert "Rahasia penalaran" not in markdown


def test_single_answer_export_rejects_user_or_missing_message() -> None:
    with pytest.raises(ValueError, match="Hanya jawaban"):
        build_chat_message_markdown("Chat", MESSAGES, "u1")
    with pytest.raises(ValueError, match="tidak ditemukan"):
        build_chat_message_markdown("Chat", MESSAGES, "missing")


def test_full_chat_export_contains_each_turn_and_excludes_thinking() -> None:
    markdown, _ = build_chat_session_markdown("Analisis", MESSAGES)

    assert "### Pengguna" in markdown
    assert "### Legal-Verse AI" in markdown
    assert "Rahasia penalaran" not in markdown


def test_docx_renderer_creates_readable_headings_and_table() -> None:
    markdown, _ = build_chat_message_markdown("Analisis Putusan", MESSAGES, "a1")
    content = docx_bytes("Analisis Putusan", markdown)

    assert content.startswith(b"PK")
    document = Document(BytesIO(content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Analisis Putusan" in text
    assert "Apa amar putusan perkara ini?" in text
    assert "Terdakwa dinyatakan bersalah." in text
    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 0).text == "Unsur"
    assert document.tables[0].cell(1, 1).text == "Terbukti"


def test_markdown_bytes_and_filename_are_download_safe() -> None:
    assert markdown_bytes("# Halo").endswith(b"\n")
    assert safe_filename("Putusan: No. 1/2025 — PN Jakarta") == "Putusan-No.-1-2025-PN-Jakarta"


def _fake_document(data: dict):
    snapshot = SimpleNamespace(exists=True, to_dict=lambda: data)
    reference = mock.Mock()
    reference.get.return_value = snapshot
    collection = mock.Mock()
    collection.document.return_value = reference
    fake_db = mock.Mock()
    fake_db.collection.return_value = collection
    return fake_db


def test_library_markdown_export_requires_owner_and_sets_download_headers() -> None:
    fake_db = _fake_document(
        {
            "user_id": "uid-123",
            "name": "Putusan 7.pdf",
            "extension": "pdf",
            "text": "[HALAMAN 1]\nMENGADILI\nPidana dua tahun.",
            "page_count": 1,
        }
    )
    request = ExportRequest(source_type="library_file", source_id="file-1", format="md")

    with mock.patch("app.controllers.export_controller.db", fake_db):
        response = export_document(request, {"uid": "uid-123"})

    assert response.media_type == "text/markdown"
    assert response.body.startswith(b"# Putusan 7")
    assert "attachment" in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"


def test_export_rejects_non_owner() -> None:
    fake_db = _fake_document({"user_id": "someone-else", "name": "private.pdf"})
    request = ExportRequest(source_type="library_file", source_id="file-1", format="md")

    with mock.patch("app.controllers.export_controller.db", fake_db):
        with pytest.raises(HTTPException) as raised:
            export_document(request, {"uid": "uid-123"})

    assert raised.value.status_code == 403
